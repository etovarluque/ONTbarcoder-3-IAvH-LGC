"""Standalone analysis: for a completed Parameter Batch run, work out exactly
which run combinations were needed to recover the taxonomically-identified
best sequences produced by Best Sequence Selection (best_seq_panel.py).

Given:
  - a Parameter Batch output folder (the "..._batch" folder, containing
    batch_run_summary.tsv, produced by batch_sweep.py), and
  - a Best Sequence Selection output folder (the "..._bestseq" folder,
    containing bestseq-<date>.tsv and bestseq-<date>_identified.fasta),

this script:
  1. Re-derives, for every sample with a taxonomic hit, exactly which of the
     N run folders reproduce that winning sequence identically (by direct
     sequence comparison against each run's consensus_filtered.fa — not by
     trusting the ";run<N>" tag in the dedup FASTA, which only records the
     FIRST run that produced a given variant, not every run that did).
  2. Reports per-run coverage (how many identified sequences a single run
     alone would reproduce).
  3. Computes a greedy minimal set of runs that together cover as many of
     the identified sequences as possible.
  4. Writes a multi-sheet Excel report with all of the above.

No Qt/PyQt import here on purpose: runnable standalone from the command
line, independent of the GUI.

Usage:
    python batch_coverage_report.py --batch-dir <..._batch folder> \
        --bestseq-dir <..._bestseq folder> [--output report.xlsx]
"""
from __future__ import annotations

import argparse
import csv
import glob
import os
import sys
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FASTA_NAME_DEFAULT = "consensus_filtered.fa"


# ---------------------------------------------------------------------------
# FASTA helpers
# ---------------------------------------------------------------------------

def read_fasta(path: str) -> List[Tuple[str, str]]:
    """Return [(header_without_'>', sequence), ...] in file order."""
    records: List[Tuple[str, str]] = []
    header = None
    seq_parts: List[str] = []
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.rstrip("\n")
            if line.startswith(">"):
                if header is not None:
                    records.append((header, "".join(seq_parts)))
                header = line[1:]
                seq_parts = []
            else:
                seq_parts.append(line.strip())
    if header is not None:
        records.append((header, "".join(seq_parts)))
    return records


def sample_of(header: str) -> str:
    """Sample id: header text before the first ';' (ONTbarcoder convention)."""
    return header.split(";", 1)[0].strip()


def parse_params(param_str: str) -> Dict[str, str]:
    """'k1=v1, k2=v2, ...' -> {'k1': 'v1', 'k2': 'v2', ...}."""
    out: Dict[str, str] = {}
    for part in param_str.split(","):
        part = part.strip()
        if not part or "=" not in part:
            continue
        k, v = part.split("=", 1)
        out[k.strip()] = v.strip()
    return out


# ---------------------------------------------------------------------------
# Input discovery / loading
# ---------------------------------------------------------------------------

def find_one(pattern: str, description: str) -> str:
    matches = [p for p in glob.glob(pattern) if not os.path.basename(p).startswith("~$")]
    if not matches:
        raise FileNotFoundError(f"No {description} found matching: {pattern}")
    matches.sort(key=os.path.getmtime, reverse=True)
    if len(matches) > 1:
        print(f"  Note: {len(matches)} {description} candidates found, "
              f"using most recent: {os.path.basename(matches[0])}", file=sys.stderr)
    return matches[0]


def load_runs(batch_dir: str, output_root: str) -> Dict[int, dict]:
    """Parse batch_run_summary.tsv -> {run_number: {folder, params, params_raw, n_consensus_filtered}}."""
    summary_path = os.path.join(batch_dir, "batch_run_summary.tsv")
    if not os.path.isfile(summary_path):
        raise FileNotFoundError(f"batch_run_summary.tsv not found in {batch_dir}")

    runs: Dict[int, dict] = {}
    with open(summary_path, encoding="utf-8") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        for row in reader:
            run_no = int(row["Run"])
            runs[run_no] = {
                "folder": os.path.join(output_root, "ont-barcoder_" + row["Folder"]),
                "params_raw": row["Parameters"],
                "params": parse_params(row["Parameters"]),
                "n_consensus_filtered": row.get("N_consensus_filtered", ""),
            }
    return runs


def load_tax_info(bestseq_tsv: str) -> Dict[str, dict]:
    with open(bestseq_tsv, encoding="utf-8") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        return {row["Sample"]: row for row in reader}


# ---------------------------------------------------------------------------
# Core analysis
# ---------------------------------------------------------------------------

def compute_coverage(
    winners: Dict[str, str], runs: Dict[int, dict], fasta_name: str
) -> Tuple[Dict[str, set], Dict[int, int]]:
    """For each sample in `winners`, find every run whose own consensus FASTA
    contains that exact sequence for that sample.

    Returns (coverage: sample -> set of run numbers, run_hit_count: run -> count).
    """
    coverage: Dict[str, set] = {s: set() for s in winners}
    run_hit_count: Dict[int, int] = {}

    for run_no, info in sorted(runs.items()):
        fa_path = os.path.join(info["folder"], fasta_name)
        hits = 0
        if os.path.isfile(fa_path):
            for header, seq in read_fasta(fa_path):
                s = sample_of(header)
                if s in winners and winners[s] == seq:
                    coverage[s].add(run_no)
                    hits += 1
        else:
            print(f"  Warning: missing {fa_path}", file=sys.stderr)
        run_hit_count[run_no] = hits

    return coverage, run_hit_count


def greedy_set_cover(
    winners: Dict[str, str], coverage: Dict[str, set], runs: Dict[int, dict]
) -> List[Tuple[int, set]]:
    """Greedy minimal set of runs covering as many samples as possible."""
    remaining = set(winners.keys())
    chosen: List[Tuple[int, set]] = []
    while remaining:
        best_r, best_cov = None, set()
        for r in runs:
            cov = {s for s in remaining if r in coverage[s]}
            if len(cov) > len(best_cov):
                best_r, best_cov = r, cov
        if best_r is None or not best_cov:
            break
        chosen.append((best_r, best_cov))
        remaining -= best_cov
    return chosen


# ---------------------------------------------------------------------------
# Excel report
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def style_header(ws, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def autofit(ws, widths: List[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(
    out_path: str,
    runs: Dict[int, dict],
    winners: Dict[str, str],
    tax_info: Dict[str, dict],
    coverage: Dict[str, set],
    run_hit_count: Dict[int, int],
    chosen_runs: List[Tuple[int, set]],
    param_keys: List[str],
) -> None:
    n_runs = len(runs)
    n_winners = len(winners)
    covered_by_set = set()
    sample_assigned_run: Dict[str, int] = {}
    for r, cov in chosen_runs:
        covered_by_set |= cov
        for s in cov:
            sample_assigned_run[s] = r
    never_covered = sorted(set(winners.keys()) - covered_by_set)

    wb = Workbook()

    # --- Sheet 1: Summary ---
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Parameter Batch coverage report: minimal set of run combinations "
               "needed to recover the taxonomically-identified best sequences"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Total run combinations tested (Parameter Batch)", n_runs])
    ws.append(["Total taxonomically-identified sequences", n_winners])
    ws.append(["Sequences covered by the greedy minimal set", len(covered_by_set)])
    ws.append(["Run combinations in the minimal set", len(chosen_runs)])
    ws.append(["Sequences not reproduced identically by any run folder", len(never_covered)])
    ws.append([])
    if run_hit_count:
        best_run = max(run_hit_count, key=lambda r: run_hit_count[r])
        worst_run = min(run_hit_count, key=lambda r: run_hit_count[r])
        ws.append(["Best single run", best_run, runs[best_run]["params_raw"], run_hit_count[best_run]])
        ws.append(["Worst single run", worst_run, runs[worst_run]["params_raw"], run_hit_count[worst_run]])
        ws.append(["Average sequences reproduced per single run",
                   round(sum(run_hit_count.values()) / n_runs, 1)])
    if never_covered:
        ws.append([])
        ws.append(["Samples not found identically in any run folder (check for a later re-run "
                   "of the conversion folders after the dedup/BLAST steps):"])
        for s in never_covered:
            ws.append([s])
    autofit(ws, [65, 14, 55, 14])

    # --- Sheet 2: Minimal set of combinations ---
    ws2 = wb.create_sheet("Minimal_run_set")
    headers = (["Order", "Run", "Folder"] + param_keys +
               ["New_sequences_added", "Cumulative_sequences", "Pct_of_total"])
    ws2.append(headers)
    style_header(ws2, len(headers))
    cum = 0
    for i, (r, cov) in enumerate(chosen_runs, start=1):
        cum += len(cov)
        p = runs[r]["params"]
        ws2.append([i, r, os.path.basename(runs[r]["folder"])] +
                    [p.get(k, "") for k in param_keys] +
                    [len(cov), cum, round(100 * cum / n_winners, 1) if n_winners else 0])
    autofit(ws2, [8, 6, 26] + [16] * len(param_keys) + [20, 20, 12])
    ws2.freeze_panes = "A2"

    # --- Sheet 3: Per-sample detail ---
    ws3 = wb.create_sheet("Per_sample_detail")
    headers3 = ["Sample", "Tax_level", "Query_Order", "Query_Family", "Query_Genus", "Query_organism",
                "Hit_Order", "Hit_Family", "Hit_Genus", "Hit_organism",
                "N_runs_reproducing_sequence", "Runs_reproducing_it",
                "Covered_by_minimal_set", "Run_assigned_in_minimal_set"]
    ws3.append(headers3)
    style_header(ws3, len(headers3))
    for s in sorted(winners):
        rs = sorted(coverage[s])
        info = tax_info.get(s, {})
        ws3.append([
            s, info.get("Tax_level", ""), info.get("Query_Order", ""), info.get("Query_Family", ""),
            info.get("Query_Genus", ""), info.get("Query_organism", ""),
            info.get("Hit_Order", ""), info.get("Hit_Family", ""), info.get("Hit_Genus", ""),
            info.get("Hit_organism", ""),
            len(rs), ",".join(map(str, rs)),
            "Yes" if s in covered_by_set else "No",
            sample_assigned_run.get(s, ""),
        ])
    autofit(ws3, [14, 10, 14, 16, 14, 20, 14, 16, 14, 20, 12, 30, 18, 14])
    ws3.freeze_panes = "A2"

    # --- Sheet 4: full per-run coverage (context) ---
    ws4 = wb.create_sheet("All_runs_coverage")
    headers4 = (["Run", "Folder"] + param_keys +
                ["N_consensus_filtered_total", f"N_of_{n_winners}_reproduced", "Pct"])
    ws4.append(headers4)
    style_header(ws4, len(headers4))
    for r in sorted(runs):
        p = runs[r]["params"]
        ws4.append([r, os.path.basename(runs[r]["folder"])] +
                    [p.get(k, "") for k in param_keys] +
                    [runs[r]["n_consensus_filtered"], run_hit_count.get(r, 0),
                     round(100 * run_hit_count.get(r, 0) / n_winners, 1) if n_winners else 0])
    autofit(ws4, [8, 26] + [16] * len(param_keys) + [22, 20, 10])
    ws4.freeze_panes = "A2"

    wb.save(out_path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--batch-dir", required=True,
                     help="Parameter Batch output folder (contains batch_run_summary.tsv)")
    ap.add_argument("--bestseq-dir", required=True,
                     help="Best Sequence Selection output folder (contains bestseq-<date>.tsv "
                          "and bestseq-<date>_identified.fasta)")
    ap.add_argument("--output-root", default=None,
                     help="Folder containing the run folders referenced by batch_run_summary.tsv "
                          "(default: parent folder of --batch-dir)")
    ap.add_argument("--fasta-name", default=FASTA_NAME_DEFAULT,
                     help=f"Per-run FASTA to compare against (default: {FASTA_NAME_DEFAULT})")
    ap.add_argument("--output", default=None,
                     help="Output .xlsx path (default: <bestseq-dir>/parameter_batch_coverage_report.xlsx)")
    args = ap.parse_args()

    batch_dir = os.path.abspath(args.batch_dir)
    bestseq_dir = os.path.abspath(args.bestseq_dir)
    output_root = os.path.abspath(args.output_root) if args.output_root else os.path.dirname(batch_dir)
    out_path = args.output or os.path.join(bestseq_dir, "parameter_batch_coverage_report.xlsx")

    print("Loading run combinations...")
    runs = load_runs(batch_dir, output_root)
    print(f"  {len(runs)} run combinations loaded")

    print("Locating Best Sequence Selection output...")
    bestseq_tsv = find_one(os.path.join(bestseq_dir, "bestseq-*.tsv"), "bestseq report .tsv")
    identified_fasta = find_one(os.path.join(bestseq_dir, "bestseq-*_identified.fasta"),
                                 "bestseq _identified.fasta")

    tax_info = load_tax_info(bestseq_tsv)
    identified = read_fasta(identified_fasta)
    winners = {sample_of(h): seq for h, seq in identified}
    print(f"  {len(winners)} taxonomically-identified sequences loaded")

    # Discover parameter keys in the order they first appear, so the report
    # adapts to whichever parameters were actually swept.
    param_keys: List[str] = []
    for info in runs.values():
        for k in info["params"]:
            if k not in param_keys:
                param_keys.append(k)

    print("Comparing each run's consensus FASTA against the winning sequences...")
    coverage, run_hit_count = compute_coverage(winners, runs, args.fasta_name)

    print("Computing greedy minimal run set...")
    chosen_runs = greedy_set_cover(winners, coverage, runs)
    covered = sum(len(c) for _r, c in chosen_runs)
    print(f"  {len(chosen_runs)} runs needed to cover {covered}/{len(winners)} sequences")

    print(f"Writing report to {out_path}")
    build_workbook(out_path, runs, winners, tax_info, coverage, run_hit_count,
                   chosen_runs, param_keys)
    print("Done.")


if __name__ == "__main__":
    main()
