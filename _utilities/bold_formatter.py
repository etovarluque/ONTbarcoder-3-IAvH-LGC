from __future__ import annotations
import os
import datetime
from collections import OrderedDict
from typing import Optional
from PyQt5 import QtCore, QtGui, QtWidgets
from .shared import *
from .shared import _get_base_dir, _tr
from .fasta_tools import _DragDropLineEdit


class _BoldFormatWorker(QtCore.QThread):
    """Apply the BLAST-results visual format to BOLD Barcode-ID xlsx exports.

    Takes a BOLD "Barcode ID" workbook (a title row, a header row, then hit
    rows grouped by Query ID and sorted by ID% descending) and produces a
    styled copy that mirrors the look of the BLAST results workbook.
    """

    progress = QtCore.pyqtSignal(str)   # short status shown in the label
    log_line = QtCore.pyqtSignal(str)   # detailed per-file result appended to the log
    finished = QtCore.pyqtSignal(list)
    error    = QtCore.pyqtSignal(str)

    # --- Palette taken from the BLAST results workbook -----------------------
    GRID_COLOR  = "FFCCCCCC"      # thin cell borders
    SEP_COLOR   = "FF9AA0A6"      # medium separator between Query ID groups
    HEADER_TEXT = "FFFFFFFF"
    FONT_SIZE   = 10.0

    QUERY_COL_NAME   = "Query ID"
    IDPCT_COL_NAME   = "ID%"
    HITRANK_COL_NAME = "Hit rank"

    # Semantic column groups by header name -> ("header_rgb", "band_rgb").
    _BLUE   = ("FF1A365D", "FFE8F1FB")
    _TEAL   = ("FF0D5E6E", "FFE8F5F6")
    _ORANGE = ("FF7C3200", "FFFEF3E8")

    COLOR_BY_HEADER = {
        "Hit rank":  _BLUE,
        "Query ID":  _BLUE,
        "PID [BIN]": _TEAL,
        "Phylum":    _ORANGE,
        "Class":     _ORANGE,
        "Order":     _ORANGE,
        "Family":    _ORANGE,
        "Subfamily": _ORANGE,
        "Tribe":     _ORANGE,
        "Genus":     _ORANGE,
        "Species":   _ORANGE,
        "Indels":    _TEAL,
        "ID%":       _TEAL,
    }

    WIDTH_BY_HEADER = {
        "Hit rank":  10,
        "Query ID":  55,
        "PID [BIN]": 16,
        "Phylum":    15,
        "Class":     15,
        "Order":     15,
        "Family":    16,
        "Subfamily": 14,
        "Tribe":     12,
        "Genus":     18,
        "Species":   28,
        "Indels":     9,
        "ID%":        9,
    }

    def __init__(self, files: list, max_hits: Optional[int], out_dir: str, parent=None):
        super().__init__(parent)
        self._files    = files
        self._max_hits = max_hits   # None = keep all hits
        self._out_dir  = out_dir
        self._stop     = False

    def stop(self):
        self._stop = True

    # ── core logic ──────────────────────────────────────────────────────────

    @classmethod
    def _read_bold(cls, path):
        """Return (headers, groups, q_idx, id_idx).

        The input has a title in row 1, headers in row 2 and data from row 3 on.
        Rows keep their original relative order (already sorted by ID% desc).
        """
        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        max_col = ws.max_column

        headers = [ws.cell(row=2, column=c).value for c in range(1, max_col + 1)]
        if cls.QUERY_COL_NAME not in headers:
            raise ValueError(f"no '{cls.QUERY_COL_NAME}' column found")
        if cls.IDPCT_COL_NAME not in headers:
            raise ValueError(f"no '{cls.IDPCT_COL_NAME}' column found")

        q_idx = headers.index(cls.QUERY_COL_NAME)
        id_idx = headers.index(cls.IDPCT_COL_NAME)

        groups = OrderedDict()
        for r in range(3, ws.max_row + 1):
            row = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
            if row[q_idx] is None and all(v is None for v in row):
                continue  # skip fully blank rows
            groups.setdefault(row[q_idx], []).append(row)

        return headers, groups, q_idx, id_idx

    @classmethod
    def _build(cls, headers, groups, q_idx, id_idx, max_hits):
        """Sort/trim each group and return (out_headers, records, reordered).

        records is a list of (row_values, hit_rank, group_index, is_first_in_group).
        A "Hit rank" column is prepended.
        """
        out_headers = [cls.HITRANK_COL_NAME] + headers
        records = []
        reordered = []

        for gi, (query, rows) in enumerate(groups.items(), start=1):
            def id_key(row):
                v = row[id_idx]
                return v if isinstance(v, (int, float)) else float("-inf")

            vals = [id_key(r) for r in rows]
            if any(a < b for a, b in zip(vals, vals[1:])):
                reordered.append(query)
            rows_sorted = sorted(rows, key=id_key, reverse=True)

            if max_hits is not None:
                rows_sorted = rows_sorted[:max_hits]

            for rank, row in enumerate(rows_sorted, start=1):
                records.append(([rank] + row, rank, gi, rank == 1))

        return out_headers, records, reordered

    @classmethod
    def _write(cls, out_headers, records, out_path):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        def _side(style, rgb):
            return Side(style=style, color=rgb)

        wb = Workbook()
        ws = wb.active
        ws.title = "Barcode ID"

        thin = _side("thin", cls.GRID_COLOR)
        border_thin = Border(top=thin, bottom=thin, left=thin, right=thin)
        border_sep = Border(top=_side("medium", cls.SEP_COLOR),
                            bottom=thin, left=thin, right=thin)

        header_fills = {}
        band_fills = {}
        for name in set(out_headers):
            hrgb, brgb = cls.COLOR_BY_HEADER.get(name, cls._TEAL)
            header_fills[name] = PatternFill("solid", fgColor=hrgb)
            band_fills[name] = PatternFill("solid", fgColor=brgb)

        center = Alignment(horizontal="center", vertical="center")

        # Header row.
        for c, name in enumerate(out_headers, start=1):
            cell = ws.cell(row=1, column=c, value=name)
            cell.fill = header_fills[name]
            cell.font = Font(bold=True, color=cls.HEADER_TEXT, size=cls.FONT_SIZE)
            cell.border = border_thin
            cell.alignment = center
        ws.row_dimensions[1].height = 22   # matches the BLAST results workbook

        idpct_col = out_headers.index(cls.IDPCT_COL_NAME) + 1

        # Data rows.
        for i, (values, rank, gi, is_first) in enumerate(records):
            r = i + 2
            banded = (gi % 2 == 1)  # odd groups colored, even groups white
            for c, name in enumerate(out_headers, start=1):
                cell = ws.cell(row=r, column=c, value=values[c - 1])
                cell.font = Font(bold=is_first, size=cls.FONT_SIZE)
                if banded:
                    cell.fill = band_fills[name]
                cell.border = border_sep if (is_first and gi > 1) else border_thin
                if name == cls.HITRANK_COL_NAME:
                    cell.alignment = center
            ws.cell(row=r, column=idpct_col).number_format = "0.00"

        # Column widths, freeze and filter.
        for c, name in enumerate(out_headers, start=1):
            ws.column_dimensions[get_column_letter(c)].width = cls.WIDTH_BY_HEADER.get(name, 15)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(out_headers))}{len(records) + 1}"

        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        wb.save(out_path)

    # ── thread entry point ─────────────────────────────────────────────────

    def run(self):
        try:
            outputs = []
            total = len(self._files)
            for i, path in enumerate(self._files, 1):
                if self._stop:
                    return
                name = os.path.basename(path)
                self.progress.emit(f"Formatting {name}…  ({i}/{total})")
                try:
                    headers, groups, q_idx, id_idx = self._read_bold(path)
                    out_headers, records, reordered = self._build(
                        headers, groups, q_idx, id_idx, self._max_hits)
                    out_path = os.path.join(
                        self._out_dir,
                        f"{os.path.splitext(name)[0]}_formatted.xlsx")
                    self._write(out_headers, records, out_path)
                    outputs.append(out_path)

                    kept = "all" if self._max_hits is None else self._max_hits
                    self.log_line.emit(
                        f"{name}: {len(groups)} Query ID(s) · "
                        f"{len(records)} row(s) (hits/group: {kept}) "
                        f"→ {os.path.basename(out_path)}")
                    if reordered:
                        self.log_line.emit(
                            f"  ⚠ {len(reordered)} group(s) were not sorted by "
                            f"ID% desc and got re-sorted")
                except Exception as exc:
                    self.log_line.emit(f"{name}: ERROR — {exc}")

            if self._stop:
                return
            if not outputs:
                self.error.emit(
                    "No files were formatted. Make sure the input is a BOLD "
                    "Barcode-ID export (title row, header row, then hits).")
                return
            self.finished.emit(outputs)
        except Exception as exc:  # pragma: no cover - defensive
            self.error.emit(str(exc))


class BoldFormatterPanel(QtWidgets.QWidget):

    def __init__(self, parent=None):
        super().__init__(parent)

        outer = QtWidgets.QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        # ── Scroll area ──
        scroll = QtWidgets.QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QtWidgets.QFrame.NoFrame)
        inner = QtWidgets.QWidget()
        self._layout = QtWidgets.QVBoxLayout(inner)
        self._layout.setContentsMargins(20, 20, 20, 20)
        self._layout.setSpacing(16)
        scroll.setWidget(inner)
        outer.addWidget(scroll, 1)

        # ── Title ──
        self._layout.addWidget(make_label("BOLD Formatter", size=19, bold=True))
        desc = make_label(
            "Take a BOLD “Barcode ID” .xlsx export and restyle it to match the "
            "BLAST results workbook: a per-group “Hit rank” column, top-N hits "
            "per Query ID, group colors, banded fills, frozen header and filters.",
            color=TEXT_SEC,
        )
        desc.setWordWrap(True)
        self._layout.addWidget(desc)

        # ── Input file (drag & drop + browse) ──
        self._layout.addWidget(self._make_section_lbl("Input file"))

        file_row = QtWidgets.QHBoxLayout()
        file_row.setSpacing(8)
        lbl_in = make_label("BOLD .xlsx:", color=TEXT_SEC)
        lbl_in.setToolTip(
            "BOLD 'Barcode ID' workbook: a title row, a header row with a "
            "'Query ID' and 'ID%' column, then hit rows grouped by Query ID."
        )
        file_row.addWidget(lbl_in)
        self._file_edit = _DragDropLineEdit(accepted_extensions=[".xlsx"])
        self._file_edit.setReadOnly(True)
        self._file_edit.setPlaceholderText("No file selected… (or drag & drop)")
        self._file_edit.textChanged.connect(self._on_file_changed)
        file_row.addWidget(self._file_edit, 1)
        self._browse_btn = QtWidgets.QPushButton("Browse…")
        self._browse_btn.setObjectName("secondary_btn")
        self._browse_btn.setFixedWidth(120)
        self._browse_btn.clicked.connect(self._browse_file)
        file_row.addWidget(self._browse_btn)
        self._layout.addLayout(file_row)

        # ── Options ──
        self._layout.addWidget(self._make_section_lbl("Options"))

        hits_row = QtWidgets.QHBoxLayout()
        hits_row.setSpacing(12)
        lbl_hits = make_label("Hits to keep per Query ID:", color=TEXT_SEC)
        hits_row.addWidget(lbl_hits)
        self._hits_spin = QtWidgets.QSpinBox()
        self._hits_spin.setRange(1, 999)
        self._hits_spin.setValue(5)
        self._hits_spin.setFixedWidth(90)
        hits_row.addWidget(self._hits_spin)
        hits_row.addStretch()
        self._layout.addLayout(hits_row)

        self._keep_all_chk = QtWidgets.QCheckBox("Keep all hits (ignore the limit above)")
        self._keep_all_chk.toggled.connect(
            lambda on: (self._hits_spin.setDisabled(on), lbl_hits.setDisabled(on)))
        self._layout.addWidget(self._keep_all_chk)

        # ── Status + progress ──
        self._status_lbl = make_label("", color=TEXT_SEC)
        self._layout.addWidget(self._status_lbl)

        self._progress_bar = QtWidgets.QProgressBar()
        self._progress_bar.setRange(0, 0)
        self._progress_bar.setFixedHeight(6)
        self._progress_bar.hide()
        self._layout.addWidget(self._progress_bar)

        self._layout.addStretch()

        # ── Operation log ──
        self._log_edit = QtWidgets.QPlainTextEdit()
        self._log_edit.setReadOnly(True)
        self._log_edit.setFixedHeight(110)
        self._log_edit.setStyleSheet(f"""
            QPlainTextEdit {{
                background: {GRAY_BG};
                border: 1px solid {GRAY_LINE};
                border-top: 1px solid {GRAY_LINE};
                border-radius: 0px;
                font-family: 'Courier New', Consolas, monospace;
                font-size: 16px;
                color: {TEXT_PRI};
                padding: 6px;
            }}
        """)
        self._log_edit.hide()
        outer.addWidget(self._log_edit)

        # ── Footer ──
        footer = QtWidgets.QWidget()
        footer.setObjectName("bold_formatter_footer")
        footer.setStyleSheet(f"""
            QWidget#bold_formatter_footer {{
                background: {GRAY_CARD};
                border-top: 1px solid {GRAY_LINE};
            }}
        """)
        fl = QtWidgets.QHBoxLayout(footer)
        fl.setContentsMargins(20, 10, 20, 10)
        fl.setSpacing(8)

        self._clear_btn = QtWidgets.QPushButton("Clear")
        self._clear_btn.setObjectName("danger_btn")
        self._clear_btn.setFixedHeight(44)
        self._clear_btn.setFixedWidth(140)
        self._clear_btn.clicked.connect(self._clear)
        fl.addWidget(self._clear_btn)

        self._open_folder_btn = QtWidgets.QPushButton("Open folder  📂")
        self._open_folder_btn.setObjectName("secondary_btn")
        self._open_folder_btn.setFixedHeight(44)
        self._open_folder_btn.hide()
        self._open_folder_btn.clicked.connect(self._open_output_folder)
        fl.addWidget(self._open_folder_btn)
        fl.addStretch()

        self._run_btn = QtWidgets.QPushButton("Format  →")
        self._run_btn.setObjectName("primary_btn")
        self._run_btn.setFixedHeight(44)
        self._run_btn.setFixedWidth(300)
        self._run_btn.clicked.connect(self._on_run_clicked)
        fl.addWidget(self._run_btn)
        self._set_run_enabled(False)
        outer.addWidget(footer)

        self._worker: Optional[_BoldFormatWorker] = None
        self._retired_workers: set = set()
        self._last_outputs: list = []

    # ── helpers ─────────────────────────────────────────────────────────────

    @staticmethod
    def _make_section_lbl(text):
        lbl = make_section_label(text)
        lbl.setStyleSheet(
            f"font-size:17px; font-weight:600; color:{BLUE}; letter-spacing:0.4px;"
        )
        return lbl

    def _set_run_enabled(self, enabled: bool):
        self._run_btn.setEnabled(enabled)
        if enabled:
            self._run_btn.setStyleSheet(
                f"QPushButton {{ background-color:{BLUE}; color:white; border:none; "
                f"border-radius:8px; padding:9px 20px; font-size:18px; font-weight:500; }}"
                f"QPushButton:hover {{ background-color:#0C4A82; }}"
            )
        else:
            self._run_btn.setStyleSheet(
                f"QPushButton {{ background-color:{GRAY_LINE}; color:{TEXT_HINT}; border:none; "
                f"border-radius:8px; padding:9px 20px; font-size:18px; font-weight:500; }}"
            )

    def _retire_worker(self):
        """Detach the current worker so a late finish can't touch the UI, and keep
        a reference until its thread exits so it is never garbage-collected while
        still running (which crashes Qt)."""
        self._retired_workers = {w for w in self._retired_workers if w.isRunning()}
        w = self._worker
        self._worker = None
        if w is None:
            return
        for sig in (w.progress, w.log_line, w.finished, w.error):
            try:
                sig.disconnect()
            except (TypeError, RuntimeError):
                pass
        if w.isRunning():
            w.stop()
            self._retired_workers.add(w)

    # ── slots ────────────────────────────────────────────────────────────────

    def _browse_file(self):
        path, _ = QtWidgets.QFileDialog.getOpenFileName(
            self, "Select BOLD Barcode-ID .xlsx", "", "Excel files (*.xlsx)"
        )
        if path:
            self._file_edit.setText(path)

    def _on_file_changed(self, text: str):
        self._set_run_enabled(bool(text.strip()))
        self._status_lbl.setText("")
        self._status_lbl.setStyleSheet("")
        self._open_folder_btn.hide()

    def _clear(self):
        self._file_edit.clear()
        self._hits_spin.setValue(5)
        self._keep_all_chk.setChecked(False)
        self._status_lbl.setText("")
        self._status_lbl.setStyleSheet("")
        self._progress_bar.hide()
        self._log_edit.clear()
        self._log_edit.hide()
        self._open_folder_btn.hide()
        self._set_run_enabled(False)

    def _on_run_clicked(self):
        path = self._file_edit.text().strip()
        if not path:
            return
        if not os.path.isfile(path):
            QtWidgets.QMessageBox.warning(
                self, "Input file", f"File not found:\n{path}")
            return

        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        folder_name = f"ont-barcoder_{ts}_bold_format"
        auto_dir = os.path.join(_get_base_dir(), "output", folder_name)

        out_dir = self._ask_output_dir(auto_dir, folder_name)
        if out_dir is None:
            return

        max_hits = None if self._keep_all_chk.isChecked() else self._hits_spin.value()

        self._progress_bar.show()
        self._status_lbl.setStyleSheet("")
        self._status_lbl.setText("Formatting…")
        self._set_run_enabled(False)
        self._log_edit.clear()
        self._log_edit.show()
        self._open_folder_btn.hide()

        self._retire_worker()
        self._worker = _BoldFormatWorker([path], max_hits, out_dir)
        self._worker.progress.connect(self._on_progress)
        self._worker.log_line.connect(self._on_log_line)
        self._worker.finished.connect(self._on_finished)
        self._worker.error.connect(self._on_error)
        self._worker.start()

    def _ask_output_dir(self, auto_dir: str, folder_name: str) -> Optional[str]:
        dlg = QtWidgets.QDialog(self)
        dlg.setWindowTitle("Output folder")
        dlg.setMinimumWidth(480)
        dlg.setStyleSheet(f"""
            QDialog {{ background-color: {GRAY_CARD}; }}
            QLabel {{ color: {TEXT_PRI}; background-color: transparent; }}
            QRadioButton {{
                color: {TEXT_PRI}; background-color: transparent;
                font-size: 15px; padding: 6px 0;
            }}
            QRadioButton::indicator {{ width: 16px; height: 16px; }}
            QPushButton {{
                border-radius: 8px; padding: 8px 20px;
                font-size: 15px; font-weight: 500;
            }}
            #dlg_ok_btn {{ background-color: {BLUE}; color: white; border: none; }}
            #dlg_ok_btn:hover {{ background-color: #0C4A82; }}
            #dlg_cancel_btn {{
                background-color: transparent; color: {BLUE};
                border: 1px solid {BLUE};
            }}
            #dlg_cancel_btn:hover {{ background-color: {BLUE_LIGHT}; }}
        """)

        vlay = QtWidgets.QVBoxLayout(dlg)
        vlay.setSpacing(16)
        vlay.setContentsMargins(24, 24, 24, 20)

        title_lbl = QtWidgets.QLabel("Where to save the results?")
        title_lbl.setStyleSheet(f"font-size:17px; font-weight:700; color:{TEXT_PRI};")
        vlay.addWidget(title_lbl)

        radio_auto = QtWidgets.QRadioButton(
            f"Automatic folder (recommended)\n  …/output/{folder_name}/"
        )
        radio_auto.setChecked(True)
        radio_custom = QtWidgets.QRadioButton("Select folder manually")
        vlay.addWidget(radio_auto)
        vlay.addWidget(radio_custom)
        vlay.addSpacing(8)

        btn_row = QtWidgets.QHBoxLayout()
        btn_row.addStretch()
        btn_cancel = QtWidgets.QPushButton("Cancel")
        btn_cancel.setObjectName("dlg_cancel_btn")
        btn_cancel.setFixedHeight(38)
        btn_ok = QtWidgets.QPushButton("Run")
        btn_ok.setObjectName("dlg_ok_btn")
        btn_ok.setFixedHeight(38)
        btn_ok.setDefault(True)
        btn_cancel.clicked.connect(dlg.reject)
        btn_ok.clicked.connect(dlg.accept)
        btn_row.addWidget(btn_cancel)
        btn_row.addSpacing(8)
        btn_row.addWidget(btn_ok)
        vlay.addLayout(btn_row)

        if dlg.exec_() != QtWidgets.QDialog.Accepted:
            return None

        if radio_auto.isChecked():
            return auto_dir
        parent_dir = QtWidgets.QFileDialog.getExistingDirectory(
            self, "Select output folder"
        )
        if not parent_dir:
            return None
        return os.path.join(parent_dir, folder_name)

    def _on_progress(self, msg: str):
        self._status_lbl.setText(msg)

    def _on_log_line(self, line: str):
        self._log_edit.appendPlainText(line)

    def _on_finished(self, outputs: list):
        self._progress_bar.hide()
        self._set_run_enabled(True)
        self._last_outputs = outputs
        self._status_lbl.setStyleSheet(f"color:{GREEN};")
        self._status_lbl.setText(f"Done — {len(outputs)} file(s) saved.")
        self._log_edit.appendPlainText(
            f"─── {len(outputs)} file(s) saved to: "
            f"{os.path.dirname(outputs[0]) if outputs else '—'}"
        )
        if outputs:
            self._open_folder_btn.show()

    def _on_error(self, msg: str):
        self._progress_bar.hide()
        self._set_run_enabled(True)
        self._status_lbl.setStyleSheet(f"color:{RED};")
        self._status_lbl.setText(f"Error: {msg}")
        self._log_edit.appendPlainText(f"ERROR: {msg}")

    def _open_output_folder(self):
        if not self._last_outputs:
            return
        folder = os.path.dirname(self._last_outputs[0])
        if os.path.isdir(folder):
            QtGui.QDesktopServices.openUrl(
                QtCore.QUrl.fromLocalFile(folder)
            )
